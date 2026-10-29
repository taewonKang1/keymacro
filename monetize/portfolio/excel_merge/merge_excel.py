# -*- coding: utf-8 -*-
"""
엑셀 여러 파일 합치기 + 요약 시트 자동 생성

폴더 안의 .xlsx / .csv 파일을 모두 읽어 하나의 엑셀로 합치고,
지정한 기준 열로 합계 요약 시트를 만듭니다.

사용 예:
    python merge_excel.py sample_input result.xlsx --group 지점 --sum 매출
"""

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill


def read_any(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "cp949"):
            try:
                return pd.read_csv(path, encoding=enc)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"인코딩을 알 수 없음: {path.name}")
    return pd.read_excel(path)


def merge_folder(folder: Path) -> pd.DataFrame:
    files = sorted(p for p in folder.iterdir()
                   if p.suffix.lower() in (".xlsx", ".csv") and not p.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"{folder} 에 .xlsx/.csv 파일이 없습니다.")
    frames = []
    for f in files:
        df = read_any(f)
        df.columns = [str(c).strip() for c in df.columns]
        df.insert(0, "원본파일", f.name)
        frames.append(df)
    merged = pd.concat(frames, ignore_index=True)
    return merged.dropna(how="all", subset=[c for c in merged.columns if c != "원본파일"])


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for i, col in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width * 1.6), 50)
    ws.freeze_panes = "A2"


def write_result(merged: pd.DataFrame, out: Path, group: str | None, sum_col: str | None):
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        merged.to_excel(xw, sheet_name="통합데이터", index=False)
        if group and sum_col:
            for col in (group, sum_col):
                if col not in merged.columns:
                    raise KeyError(f"열 '{col}' 이 없습니다. 사용 가능한 열: {list(merged.columns)}")
            values = pd.to_numeric(merged[sum_col], errors="coerce")
            summary = (merged.assign(**{sum_col: values})
                       .groupby(group, as_index=False)[sum_col].sum()
                       .sort_values(sum_col, ascending=False))
            summary.loc[len(summary)] = ["합계", summary[sum_col].sum()]
            summary.to_excel(xw, sheet_name="요약", index=False)
        for ws in xw.book.worksheets:
            style_sheet(ws)


def main(argv=None):
    ap = argparse.ArgumentParser(description="폴더 내 엑셀/CSV 파일 합치기")
    ap.add_argument("folder", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("--group", help="요약 기준 열 (예: 지점)")
    ap.add_argument("--sum", dest="sum_col", help="합계를 낼 열 (예: 매출)")
    args = ap.parse_args(argv)

    merged = merge_folder(args.folder)
    write_result(merged, args.output, args.group, args.sum_col)
    print(f"완료: {len(merged)}행 → {args.output}")


if __name__ == "__main__":
    sys.exit(main())
